# -*- coding: utf-8 -*-
"""Генерация 100 персональных логинов/паролей:
users.json (для встраивания в HTML) + таблица XLSX для разработчика."""
import json
import random

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

random.seed(20261011)
ALPHA = "abcdefghjkmnpqrstuvwxyz23456789"  # без 0/O/1/l/i

users = {}
for i in range(1, 101):
    login = "user%03d" % i
    while True:
        pw = "".join(random.choice(ALPHA) for _ in range(8))
        if pw not in users.values():
            break
    users[login] = {"pass": pw, "max": 5}

json.dump(users, open("/home/user/arena_manual/users.json", "w"),
          indent=1, ensure_ascii=False)

# ---------------- XLSX ----------------
wb = Workbook()
ws = wb.active
ws.title = "Пользователи (100)"

hdr_fill = PatternFill("solid", fgColor="1F3864")
hdr_font = Font(color="FFFFFF", bold=True, size=11)
thin = Side(style="thin", color="B9C4DA")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

ws.append(["№", "Логин", "Пароль", "Просмотров", "Статус"])
for c in ws[1]:
    c.fill, c.font, c.border = hdr_fill, hdr_font, border
    c.alignment = Alignment(horizontal="center")

for i, (login, rec) in enumerate(users.items(), 1):
    ws.append([i, login, rec["pass"], rec["max"], "выдан / не выдан"])
    for c in ws[ws.max_row]:
        c.border = border
    ws.cell(ws.max_row, 2).font = Font(name="Consolas")
    ws.cell(ws.max_row, 3).font = Font(name="Consolas")

ws.column_dimensions["A"].width = 6
ws.column_dimensions["B"].width = 12
ws.column_dimensions["C"].width = 12
ws.column_dimensions["D"].width = 13
ws.column_dimensions["E"].width = 18
ws.freeze_panes = "A2"

ws2 = wb.create_sheet("Служебные")
ws2.append(["Роль", "Логин", "Пароль", "Просмотров"])
for c in ws2[1]:
    c.fill, c.font = hdr_fill, hdr_font
ws2.append(["Пользователь (резервный общий)", "reader", "12345678", 5])
ws2.append(["Владелец (без лимита)", "developer", "turbo2*(+", "без лимита"])
ws2.column_dimensions["A"].width = 32
ws2.column_dimensions["B"].width = 12
ws2.column_dimensions["C"].width = 12
ws2.column_dimensions["D"].width = 13

wb.save("/home/user/Таблица_логинов_и_паролей.xlsx")
print("OK: 100 users -> users.json + Таблица_логинов_и_паролей.xlsx")
