# -*- coding: utf-8 -*-
"""Краткая презентация по инструкции AI Arena (Agent Mode) — только русский."""
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

INK = RGBColor(0x11, 0x18, 0x27)
MUT = RGBColor(0x4B, 0x55, 0x63)
ACC = RGBColor(0x1D, 0x4E, 0xD8)
ACC_SOFT = RGBColor(0xEF, 0xF4, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0x99, 0x1B, 0x1B)
RED_SOFT = RGBColor(0xFE, 0xF2, 0xF2)
GRN = RGBColor(0x06, 0x5F, 0x46)
GRN_SOFT = RGBColor(0xEC, 0xFD, 0xF5)

SH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "shots")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def txt(slide, x, y, w, h, size, color, lines, bold=False, align=PP_ALIGN.LEFT,
        space_after=6, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(ln, tuple):
            run_text, run_bold, run_color, run_size = ln
        else:
            run_text, run_bold, run_color, run_size = ln, bold, color, size
        r = p.add_run()
        r.text = run_text
        r.font.size = Pt(run_size)
        r.font.bold = run_bold
        r.font.color.rgb = run_color
        r.font.name = "Calibri"
        p.space_after = Pt(space_after)
        p.alignment = align
    return tb


def bar(slide, x, y, w, h, color):
    s = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))  # rectangle
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def header(slide, kicker, title):
    bar(slide, 0, 0, 13.333, 0.09, ACC)
    txt(slide, 0.7, 0.35, 12, 0.4, 14, ACC, [kicker], bold=True)
    txt(slide, 0.7, 0.72, 12, 1.0, 32, INK, [title], bold=True)
    bar(slide, 0.72, 1.42, 2.2, 0.05, ACC)
    txt(slide, 12.0, 7.05, 1.0, 0.35, 11, MUT,
        [("AI Arena · Agent Mode", False, MUT, 11)], align=PP_ALIGN.RIGHT)


def bullets(slide, x, y, w, items, size=18):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(5.2))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        head, rest = it
        r = p.add_run()
        r.text = "▪  " + head
        r.font.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = INK
        r.font.name = "Calibri"
        if rest:
            r2 = p.add_run()
            r2.text = " — " + rest
            r2.font.size = Pt(size)
            r2.font.color.rgb = MUT
            r2.font.name = "Calibri"
        p.space_after = Pt(12)
    return tb


def img(slide, path, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(os.path.join(SH, path), Inches(x), Inches(y), **kw)


# ---------------------------------------------------------------- 1 титул
s = prs.slides.add_slide(BLANK)
bar(s, 0, 0, 13.333, 7.5, ACC_SOFT)
bar(s, 0, 0, 13.333, 0.18, ACC)
txt(s, 0.9, 1.5, 11.5, 0.5, 16, ACC, ["КРАТКАЯ ПРЕЗЕНТАЦИЯ · ПО МАТЕРИАЛАМ ИНСТРУКЦИИ ВЕРСИИ 1.7"], bold=True)
txt(s, 0.9, 2.1, 11.5, 1.6, 44, INK, ["AI Arena: возможности Agent Mode"], bold=True)
txt(s, 0.9, 3.3, 11.0, 1.2, 22, MUT,
    ["Что умеет агент, как устроены инструменты и подключения,",
     "и какой результат вы получаете в конце."])
txt(s, 0.9, 5.2, 11.0, 0.9, 14, MUT,
    ["Источники: официальный справочный центр help.arena.ai, блог Arena,",
     "документация docs.github.com · 8 октября 2026"])
img(s, "arena_workspace.png", 8.6, 4.6, w=4.3)

# ---------------------------------------------------------------- 2 что это
s = prs.slides.add_slide(BLANK)
header(s, "01 · ЧТО ТАКОЕ AGENT MODE", "Одна модель делает задачу от начала и до конца")
bullets(s, 0.7, 1.8, 7.3, [
    ("Не чат, а рабочий процесс", "вы формулируете цель один раз — агент сам составляет план и выполняет его шагами"),
    ("Инструменты из коробки", "поиск в сети, песочница Linux, файлы, генерация картинок"),
    ("Контекст сохраняется", "правки «переделай таблицу», «добавь раздел» продолжаются с того же места"),
    ("Обратная связь в конце", "Yes / No / Keep working — оценка задачи и продолжение работы"),
])
img(s, "arena_feedback.png", 8.5, 2.0, w=4.3)
txt(s, 8.5, 3.5, 4.3, 0.8, 12, MUT, ["Финал задачи: оценка результата или «продолжай работать»."])

# ---------------------------------------------------------------- 3 инструменты
s = prs.slides.add_slide(BLANK)
header(s, "02 · ИНСТРУМЕНТЫ АГЕНТА", "Что агент делает руками")
bullets(s, 0.7, 1.8, 12.0, [
    ("Поиск и чтение страниц", "находит актуальные данные в интернете, читает сайты и документы"),
    ("Песочница bash", "ставит пакеты, запускает скрипты, тесты и серверы, проверяет и чинит ошибки"),
    ("Файлы", "создаёт и правит Word-совместимые, Excel-совместимые, Markdown, CSV, HTML-результаты"),
    ("Генерация изображений", "иллюстрации и макеты по запросу"),
    ("Уточняющие вопросы", "если задача неоднозначна — спросит, а не додумает"),
    ("Загрузка ваших файлов", "PNG, PDF, TXT, MD, CSV, HTML, CSS, JS, JSON — кнопка Add files"),
])

# ---------------------------------------------------------------- 4 workspace
s = prs.slides.add_slide(BLANK)
header(s, "03 · РАБОЧЕЕ ПРОСТРАНСТВО", "Всё, что создано, — перед глазами")
bullets(s, 0.7, 1.8, 6.6, [
    ("Workspace", "панель справа: файлы сессии, предпросмотр, скачивание"),
    ("Diff", "изменения построчно: последний ход или вся ветка"),
    ("Checks", "статус автопроверок: проходит / упало / в процессе"),
    ("Превью", "живой просмотр запущенного в песочнице сервера"),
    ("Без GitHub — ZIP", "кнопка Download скачивает всё архивом"),
])
img(s, "arena_workspace.png", 7.7, 1.9, w=5.1)

# ---------------------------------------------------------------- 5 github
s = prs.slides.add_slide(BLANK)
header(s, "04 · ГITHUB", "Результат уходит в ваш репозиторий")
bullets(s, 0.7, 1.8, 6.6, [
    ("Подключение в один клик", "OAuth, выбор репозиториев — Manage repositories"),
    ("Агент работает в копии", "клон в песочнице, коммиты и Pull Request — сам"),
    ("Ревью и мерж — ваши", "Diff и Checks помогают принять работу"),
    ("Хранилище документов", "папки shablony / activity / tp, образцы Word и Excel"),
    ("Загрузка образцов", "Add file → Upload files прямо в браузере"),
])
img(s, "gh_merge.png", 7.8, 2.0, w=4.9)
txt(s, 7.8, 4.6, 4.9, 0.8, 12, MUT, ["Pull Request: зелёная кнопка Merge — результат в основной ветке."])

# ---------------------------------------------------------------- 6 коннекторы
s = prs.slides.add_slide(BLANK)
header(s, "05 · ПОДКЛЮЧЕНИЯ", "Сервисы, к которым агент дотягивается")
bullets(s, 0.7, 1.8, 12.0, [
    ("GitHub", "код, репозитории, Pull Request — главная связка для документов и проектов"),
    ("Google и другие", "документы, файлы, почта, календарь — набор зависит от аккаунта"),
    ("Тумблер на чат", "подключение включается отдельно для каждого разговора"),
    ("Действия от вашего имени", "агент видит и делает только то, что разрешено вам"),
    ("Отключение обратимо", "выключили и включили снова — достаточно реавторизации"),
])
img(s, "arena_toggle.png", 8.6, 4.3, w=4.2)

# ---------------------------------------------------------------- 7 результат
s = prs.slides.add_slide(BLANK)
header(s, "06 · РЕЗУЛЬТАТ", "Что вы получаете в конце")
bullets(s, 0.7, 1.8, 12.0, [
    ("Готовые документы", "отчёты, справки, письма по вашим образцам — Word и Excel"),
    ("Таблицы и аналитика", "агрегаты, топ-10, поиск дублей и пропусков, графики"),
    ("Прототипы", "сайт или приложение с живым превью прямо из песочницы"),
    ("Разбор входящих файлов", "выжимки из PDF, сравнительные таблицы из нескольких документов"),
    ("Доставка", "Pull Request в репозиторий или ZIP из Workspace"),
    ("Источники", "ссылки и даты рядом с утверждениями — по запросу"),
])

# ---------------------------------------------------------------- 8 правила
s = prs.slides.add_slide(BLANK)
header(s, "07 · ПРАВИЛА БЕЗ ОПАСНОСТИ", "Три минуты, которые экономят час")
bullets(s, 0.7, 1.8, 12.0, [
    ("Забирайте результат", "всё, что осталось в песочнице, нужно скачать или отправить в репозиторий"),
    ("После мержа пуша нет", "смержили или закрыли PR — сессия больше не отправляет файлы"),
    ("Один PR на сессию", "крупную работу разбивайте на несколько чатов"),
    ("Секреты — не в промпт", "токены и пароли не вставляйте в текст задачи"),
    ("Лимиты", "Battle Mode в Text/Code Arena не тратит дневной баланс Agent Mode"),
])
bar(s, 0.7, 5.6, 11.9, 1.15, RED_SOFT)
txt(s, 0.95, 5.75, 11.4, 0.9, 15, RED,
    [("Главная ловушка: ", True, RED, 15),
     ("файлы, созданные после мержа PR, нельзя ни запушить, ни скачать, ни перенести. "
      "Сначала убедитесь, что всё отправлено, — потом мержите.", False, RED, 15)])

# ---------------------------------------------------------------- 9 старт
s = prs.slides.add_slide(BLANK)
header(s, "08 · КАК НАЧАТЬ", "Четыре шага до первого результата")
bullets(s, 0.7, 1.8, 6.6, [
    ("1. Откройте arena.ai/agent", "или переключите режим в меню слева сверху"),
    ("2. Опишите задачу простыми словами", "«сделай отчёт по образцу, вот шаблон»"),
    ("3. Прикрепите файлы", "Add files — образцы, документы, данные"),
    ("4. Нажмите стрелку", "и принимайте работу: Yes / No / Keep working"),
])
img(s, "arena_chatbar.png", 7.8, 2.1, w=4.9)
bar(s, 0.7, 5.5, 11.9, 1.1, GRN_SOFT)
txt(s, 0.95, 5.65, 11.4, 0.9, 16, GRN,
    [("Подробности — в «Инструкции по работе с AI Arena» (версия 1.7): ", True, GRN, 16),
     ("16 глав, пошаговый раздел для новичка, кликабельные ссылки на Arena и GitHub.", False, GRN, 16)])

out = "/home/user/Презентация_AI_Arena_Agent_Mode.pptx"
prs.save(out)
print("OK:", out, os.path.getsize(out), "bytes;", len(prs.slides._sldIdLst), "slides")
