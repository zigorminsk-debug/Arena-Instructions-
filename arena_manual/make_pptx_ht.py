# -*- coding: utf-8 -*-
"""Краткая презентация AI Arena (Agent Mode) — high-tech дизайн, только русский."""
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BG = RGBColor(0x07, 0x0B, 0x14)
PANEL = RGBColor(0x0D, 0x15, 0x26)
PANEL2 = RGBColor(0x10, 0x1B, 0x31)
CYAN = RGBColor(0x22, 0xD3, 0xEE)
BLUE = RGBColor(0x3B, 0x82, 0xF6)
VIOLET = RGBColor(0x8B, 0x5C, 0xF6)
WHITE = RGBColor(0xE5, 0xE7, 0xEB)
MUT = RGBColor(0x94, 0xA3, 0xB8)
DIM = RGBColor(0x47, 0x55, 0x69)
GREEN = RGBColor(0x34, 0xD3, 0x99)
RED = RGBColor(0xF8, 0x71, 0x71)
AMBER = RGBColor(0xFB, 0xBF, 0x24)

MONO = "Consolas"
BODY = "Segoe UI"

SH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "shots")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
N = [0]


def rect(s, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    return sp


def txt(s, x, y, w, h, size, color, lines, bold=False, mono=False, align=PP_ALIGN.LEFT,
        spacing=6, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        t, b, c, sz, mn = (ln + (None, None, None))[:5] if isinstance(ln, tuple) else (ln, bold, color, size, mono)
        b = b if b is not None else bold
        c = c if c is not None else color
        sz = sz if sz is not None else size
        mn = mn if mn is not None else mono
        r = p.add_run()
        r.text = t
        r.font.size = Pt(sz)
        r.font.bold = b
        r.font.color.rgb = c
        r.font.name = MONO if mn else BODY
        p.space_after = Pt(spacing)
        p.alignment = align
    return tb


def decor(s):
    N[0] += 1
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    # верхняя неоновая линия
    rect(s, 0, 0, 13.333, 0.045, fill=CYAN)
    rect(s, 0, 0.045, 4.5, 0.02, fill=BLUE)
    # уголки-скобки
    for hx, hy, vx in [(0.35, 0.35, 0.35), (12.70, 0.35, 12.96), (0.35, 6.93, 0.35), (12.70, 6.93, 12.96)]:
        rect(s, hx, hy, 0.28, 0.02, fill=DIM)
        rect(s, vx, hy if hy < 1 else hy - 0.28, 0.02, 0.28, fill=DIM)
    # точечная сетка справа сверху
    for r in range(3):
        for c in range(8):
            rect(s, 10.2 + c * 0.28, 0.30 + r * 0.22, 0.035, 0.035, fill=DIM,
                 shape=MSO_SHAPE.OVAL)
    # футер
    txt(s, 0.55, 7.06, 7, 0.3, 10, DIM, ["ARENA.AI // AGENT_MODE // INSTRUKCIJA v1.7"], mono=True)
    txt(s, 11.6, 7.06, 1.2, 0.3, 10, CYAN, ["%02d / 09" % N[0]], mono=True, align=PP_ALIGN.RIGHT)
    # строка разработчика с кликабельными контактами
    tb = s.shapes.add_textbox(Inches(0.55), Inches(7.22), Inches(10.8), Inches(0.28))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]

    def run(t, color, link=None):
        r = p.add_run()
        r.text = t
        r.font.size = Pt(9)
        r.font.color.rgb = color
        r.font.name = MONO
        if link:
            r.hyperlink.address = link

    run("разработчик: Захаревич Игорь Владимирович", MUT)
    run("  ·  ", DIM)
    run("ziv@csl.by", CYAN, "mailto:ziv@csl.by")
    run("  ·  ", DIM)
    run("csl.by", CYAN, "http://csl.by")
    run("  ·  ", DIM)
    run("+375 29 337-14-12", CYAN, "tel:+375293371412")


def header(s, code, kicker, title):
    decor(s)
    txt(s, 0.7, 0.42, 12, 0.35, 13, CYAN, ["// %s · %s" % (code, kicker)], mono=True, bold=True)
    txt(s, 0.7, 0.82, 12.0, 0.8, 30, WHITE, [title], bold=True)
    rect(s, 0.72, 1.52, 1.7, 0.045, fill=CYAN)
    rect(s, 2.5, 1.52, 0.45, 0.045, fill=BLUE)
    rect(s, 3.03, 1.52, 0.16, 0.045, fill=VIOLET)


def panel(s, x, y, w, h, line=PANEL2):
    return rect(s, x, y, w, h, fill=PANEL, line=line, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)


def bullets(s, x, y, w, items, size=16, boxed=False):
    if boxed:
        panel(s, x - 0.25, y - 0.25, w + 0.5, 0.55 * len(items) + 0.55, line=PANEL2)
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(5.4))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r0 = p.add_run()
        r0.text = "▸ "
        r0.font.size = Pt(size)
        r0.font.color.rgb = CYAN
        r0.font.bold = True
        r0.font.name = MONO
        head, rest = it
        r = p.add_run()
        r.text = head
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = WHITE
        r.font.name = BODY
        if rest:
            r2 = p.add_run()
            r2.text = " — " + rest
            r2.font.size = Pt(size - 1)
            r2.font.color.rgb = MUT
            r2.font.name = BODY
        p.space_after = Pt(13)
    return tb


RATIO = {"arena_workspace.png": 322 / 1059.0, "arena_feedback.png": 228 / 787.0,
         "arena_chatbar.png": 173 / 698.0, "arena_toggle.png": 486 / 888.0,
         "gh_merge.png": 576 / 934.0}


def screen(s, path, x, y, w, tag=None):
    h = w * RATIO[path]
    rect(s, x - 0.09, y - 0.09, w + 0.18, h + 0.18, fill=PANEL, line=BLUE, lw=1.0,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, x - 0.09, y - 0.09, w + 0.18, 0.035, fill=CYAN)
    s.shapes.add_picture(os.path.join(SH, path), Inches(x), Inches(y), width=Inches(w))
    if tag:
        txt(s, x - 0.09, y + h + 0.14, w + 0.18, 0.3, 10, DIM, [tag], mono=True, align=PP_ALIGN.CENTER)


# ---------------------------------------------------------------- 1 титул
s = prs.slides.add_slide(BLANK)
decor(s)
rect(s, 0.7, 1.15, 0.12, 2.6, fill=CYAN)
txt(s, 1.05, 1.2, 8.5, 0.4, 14, CYAN, ["[ ARENA_AI // AGENT_MODE // OVERVIEW ]"], mono=True, bold=True)
txt(s, 1.05, 1.7, 8.5, 0.9, 44, WHITE, [("AI Arena:", True, WHITE, 44, False)], spacing=0)
txt(s, 1.05, 2.55, 7.6, 0.8, 34, CYAN, [("возможности Agent Mode", True, CYAN, 34, False)], spacing=0)
txt(s, 1.05, 3.5, 7.5, 1.0, 19, MUT,
    ["Что умеет агент, как устроены инструменты и подключения —", "и какой результат вы получаете в конце."])
panel(s, 1.05, 4.5, 7.6, 1.5, line=PANEL2)
txt(s, 1.3, 4.7, 7.2, 1.2, 13, GREEN, [
    ("> agent.run(task)", True, GREEN, 14, True),
    ("  tools: [web_search, bash, files, images, github]", False, MUT, 12, True),
    ("  result: docs // tables // prototypes // PR", False, MUT, 12, True)], spacing=4)
txt(s, 1.05, 6.3, 9.5, 0.5, 11, DIM,
    ["Источники: help.arena.ai · блог Arena · docs.github.com · 8 октября 2026"], mono=True)
screen(s, "arena_workspace.png", 8.9, 2.1, 3.9, tag="workspace://session_files")

# ---------------------------------------------------------------- 2 что это
s = prs.slides.add_slide(BLANK)
header(s, "01", "ЧТО ТАКОЕ AGENT MODE", "Одна модель делает задачу от начала и до конца")
bullets(s, 0.7, 1.95, 7.2, [
    ("Не чат, а рабочий процесс", "цель формулируете один раз — агент сам строит план и выполняет шагами"),
    ("Инструменты из коробки", "поиск в сети, песочница Linux, файлы, генерация картинок"),
    ("Контекст сохраняется", "«переделай таблицу», «добавь раздел» продолжаются с того же места"),
    ("Обратная связь в конце", "Yes / No / Keep working — оценка и продолжение работы"),
])
screen(s, "arena_feedback.png", 8.5, 2.2, 4.3, tag="feedback://was_this_task_successful")

# ---------------------------------------------------------------- 3 инструменты
s = prs.slides.add_slide(BLANK)
header(s, "02", "ИНСТРУМЕНТЫ АГЕНТА", "Что агент делает руками")
bullets(s, 0.7, 1.95, 12.0, [
    ("Поиск и чтение страниц", "актуальные данные из интернета, сайты и документы"),
    ("Песочница bash", "пакеты, скрипты, тесты, серверы; проверяет и чинит ошибки"),
    ("Файлы", "Word- и Excel-совместимые отчёты, Markdown, CSV, HTML"),
    ("Генерация изображений", "иллюстрации и макеты по запросу"),
    ("Уточняющие вопросы", "если задача неоднозначна — спросит, а не додумает"),
    ("Загрузка ваших файлов", "Add files: PNG, PDF, TXT, MD, CSV, HTML, CSS, JS, JSON"),
])

# ---------------------------------------------------------------- 4 workspace
s = prs.slides.add_slide(BLANK)
header(s, "03", "РАБОЧЕЕ ПРОСТРАНСТВО", "Всё, что создано, — перед глазами")
bullets(s, 0.7, 1.95, 6.6, [
    ("Workspace", "панель справа: файлы сессии, предпросмотр, скачивание"),
    ("Diff", "изменения построчно: последний ход или вся ветка"),
    ("Checks", "статус автопроверок: проходит / упало / в процессе"),
    ("Превью", "живой просмотр сервера, запущенного в песочнице"),
    ("Без GitHub — ZIP", "кнопка Download скачивает всё архивом"),
])
screen(s, "arena_workspace.png", 7.8, 2.1, 5.0, tag="workspace://arena_expert_*")

# ---------------------------------------------------------------- 5 github
s = prs.slides.add_slide(BLANK)
header(s, "04", "GITHUB", "Результат уходит в ваш репозиторий")
bullets(s, 0.7, 1.95, 6.8, [
    ("Подключение в один клик", "OAuth и выбор репозиториев — Manage repositories"),
    ("Агент работает в копии", "клон в песочнице; коммиты и Pull Request — сам"),
    ("Ревью и мерж — ваши", "Diff и Checks помогают принять работу"),
    ("Хранилище документов", "папки shablony / activity / tp, образцы Word и Excel"),
    ("Загрузка образцов", "Add file → Upload files прямо в браузере"),
])
screen(s, "gh_merge.png", 8.1, 2.0, 4.7, tag="github://pull_request → merge")

# ---------------------------------------------------------------- 6 коннекторы
s = prs.slides.add_slide(BLANK)
header(s, "05", "ПОДКЛЮЧЕНИЯ", "Сервисы, к которым агент дотягивается")
bullets(s, 0.7, 1.95, 7.4, [
    ("GitHub", "код, репозитории, PR — главная связка для документов"),
    ("Google и другие", "документы, файлы, почта, календарь — набор по аккаунту"),
    ("Тумблер на чат", "подключение включается отдельно для разговора"),
    ("От вашего имени", "агент видит и делает только разрешённое вам"),
    ("Отключение обратимо", "повторное включение — реавторизация, ресурсы сохраняются"),
])
screen(s, "arena_toggle.png", 8.6, 2.1, 4.2, tag="connections://github = ON")

# ---------------------------------------------------------------- 7 результат
s = prs.slides.add_slide(BLANK)
header(s, "06", "РЕЗУЛЬТАТ", "Что вы получаете в конце")
bullets(s, 0.7, 1.95, 12.0, [
    ("Готовые документы", "отчёты, справки, письма по вашим образцам — Word и Excel"),
    ("Таблицы и аналитика", "агрегаты, топ-10, дубли и пропуски, графики"),
    ("Прототипы", "сайт или приложение с живым превью из песочницы"),
    ("Разбор входящих файлов", "выжимки из PDF, сравнительные таблицы документов"),
    ("Доставка", "Pull Request в репозиторий или ZIP из Workspace"),
    ("Источники", "ссылки и даты рядом с утверждениями — по запросу"),
])

# ---------------------------------------------------------------- 8 правила
s = prs.slides.add_slide(BLANK)
header(s, "07", "ПРАВИЛА БЕЗ ОПАСНОСТИ", "Три минуты, которые экономят час")
bullets(s, 0.7, 1.95, 12.0, [
    ("Забирайте результат", "оставшееся в песочнице — скачать или отправить в репозиторий"),
    ("После мержа пуша нет", "смержили или закрыли PR — сессия больше не отправляет файлы"),
    ("Один PR на сессию", "крупную работу разбивайте на несколько чатов"),
    ("Секреты — не в промпт", "токены и пароли не вставляйте в текст задачи"),
    ("Лимиты", "Battle Mode в Text/Code Arena не тратит дневной баланс Agent Mode"),
])
rect(s, 0.7, 5.55, 11.93, 1.15, fill=RGBColor(0x1A, 0x0B, 0x10), line=RED, lw=1.2,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
txt(s, 0.95, 5.72, 11.4, 0.9, 14, RED, [
    ("[!] ГЛАВНАЯ ЛОВУШКА: ", True, RED, 14, True),
    ("файлы, созданные после мержа PR, нельзя ни запушить, ни скачать, ни перенести. ", False, WHITE, 14, False),
    ("Сначала убедитесь, что всё отправлено, — потом мержите.", False, MUT, 13, False)], spacing=2)

# ---------------------------------------------------------------- 9 старт
s = prs.slides.add_slide(BLANK)
header(s, "08", "КАК НАЧАТЬ", "Четыре шага до первого результата")
bullets(s, 0.7, 1.95, 6.8, [
    ("1 · Откройте arena.ai/agent", "или переключите режим в меню слева сверху"),
    ("2 · Опишите задачу простыми словами", "«сделай отчёт по образцу, вот шаблон»"),
    ("3 · Прикрепите файлы", "Add files — образцы, документы, данные"),
    ("4 · Нажмите стрелку", "и принимайте работу: Yes / No / Keep working"),
])
screen(s, "arena_chatbar.png", 8.1, 2.2, 4.7, tag="input://prompt + add_files → run")
rect(s, 0.7, 5.5, 11.93, 1.1, fill=RGBColor(0x07, 0x1A, 0x14), line=GREEN, lw=1.2,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
txt(s, 0.95, 5.68, 11.4, 0.8, 14, GREEN, [
    ("[OK] ", True, GREEN, 14, True),
    ("Подробности — в «Инструкции по работе с AI Arena» (v1.7): 16 глав, ", False, WHITE, 14, False),
    ("пошаговый раздел для новичка, кликабельные ссылки на Arena и GitHub.", False, MUT, 13, False)], spacing=2)

out = "/home/user/Презентация_AI_Arena_Agent_Mode.pptx"
prs.save(out)
print("OK:", out, os.path.getsize(out), "bytes;", len(prs.slides._sldIdLst), "slides")
